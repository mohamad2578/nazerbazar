"""کارتابل تایید نرخ و بارگذاری گروهی نرخ‌های مصوب (اداره صمت)."""
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import serializers
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from apps.core.pagination import Pagination
from apps.core.permissions import role_permission
from apps.core.scoping import scoped
from apps.core.utils import to_en_digits

from .models import OfficialPrice, Product
from .services import PRICE_REVIEWER_ROLES, set_official_price

Reviewers = role_permission(*[r for r in PRICE_REVIEWER_ROLES if r != "admin"])
PriceLoaders = role_permission("samt", "governorate")


class PendingPriceSerializer(serializers.ModelSerializer):
    product_name = serializers.CharField(source="product.name", read_only=True)
    unit_display = serializers.CharField(source="product.get_unit_display", read_only=True)
    union_name = serializers.CharField(source="product.union.name", read_only=True)
    county_name = serializers.CharField(source="product.union.chamber.county.name", read_only=True)
    current_price = serializers.IntegerField(source="product.current_price", read_only=True)
    set_by_name = serializers.CharField(source="set_by.full_name", read_only=True, default="")
    status_display = serializers.CharField(source="get_status_display", read_only=True)

    class Meta:
        model = OfficialPrice
        fields = [
            "id", "product", "product_name", "unit_display", "union_name", "county_name",
            "price", "previous_price", "current_price", "max_discount_percent", "note",
            "status", "status_display", "set_by_name", "created_at", "review_note", "reviewed_at",
        ]


@api_view(["GET"])
@permission_classes([Reviewers])
def pending_prices(request):
    """نرخ‌هایی که اتحادیه‌ها ثبت کرده‌اند و منتظر تایید هستند."""
    qs = scoped(
        OfficialPrice.objects.select_related("product__union__chamber__county", "set_by"),
        request.user,
        OfficialPrice.SCOPE,
    ).filter(status=request.query_params.get("status") or OfficialPrice.Status.PENDING)
    paginator = Pagination()
    page = paginator.paginate_queryset(qs, request)
    return paginator.get_paginated_response(PendingPriceSerializer(page, many=True).data)


@api_view(["POST"])
@permission_classes([Reviewers])
def review_price(request, pk):
    from .services import review_price as do_review

    record = get_object_or_404(
        scoped(OfficialPrice.objects.select_related("product__union"), request.user, OfficialPrice.SCOPE), pk=pk
    )
    approve = request.data.get("approve") not in (False, "false", "False", 0, "0")
    do_review(record, request.user, approve, (request.data.get("note") or "").strip())
    return Response(PendingPriceSerializer(record).data)


# ─────────────────── بارگذاری گروهی نرخ‌ها (اکسل) ───────────────────

COLUMNS = ["شناسه کالا", "نام کالا", "اتحادیه", "واحد", "نرخ فعلی (ریال)", "نرخ جدید (ریال)", "حداکثر تخفیف (٪)", "توضیح"]


def _scoped_products(user):
    return scoped(
        Product.objects.filter(is_active=True).select_related("union__chamber__county"), user, Product.SCOPE
    ).order_by("union__name", "name")


@api_view(["GET"])
@permission_classes([PriceLoaders])
def bulk_template(request):
    """فایل اکسل آماده، شامل کالاهای حوزه کاربر و نرخ فعلی آن‌ها."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "نرخ مصوب"
    ws.sheet_view.rightToLeft = True
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for p in _scoped_products(request.user):
        ws.append([p.pk, p.name, p.union.name, p.get_unit_display(), p.current_price, "", p.max_discount_percent, ""])
    for i, _ in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 20
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="prices-{timezone.localdate()}.xlsx"'
    return resp


def _parse_int(value):
    if value in (None, ""):
        return None
    try:
        return int(float(to_en_digits(str(value)).replace(",", "").replace("٬", "").strip()))
    except ValueError:
        return None


def _read_rows(file):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, data_only=True)
    except Exception:  # noqa: BLE001
        raise ValidationError({"file": "فایل اکسل قابل خواندن نیست."})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise ValidationError({"file": "فایل خالی است."})
    return rows


@api_view(["POST"])
@permission_classes([PriceLoaders])
@parser_classes([MultiPartParser])
def bulk_upload(request):
    """بارگذاری گروهی نرخ مصوب از فایل اکسل.

    با `dry_run` فقط پیش‌نمایش تغییرات برگردانده می‌شود و چیزی ذخیره نمی‌گردد.
    """
    f = request.FILES.get("file")
    if not f:
        raise ValidationError({"file": "فایل اکسل را انتخاب کنید."})
    dry_run = request.data.get("dry_run") in (True, "true", "True", 1, "1")

    allowed = {p.pk: p for p in _scoped_products(request.user)}
    changes, errors, skipped = [], [], 0

    for i, row in enumerate(_read_rows(f), start=2):
        pid = _parse_int(row[0] if row else None)
        new_price = _parse_int(row[5] if len(row) > 5 else None)
        if pid is None or new_price is None:
            skipped += 1
            continue
        product = allowed.get(pid)
        if not product:
            errors.append({"row": i, "error": f"کالای با شناسه {pid} در حوزه شما نیست."})
            continue
        if new_price <= 0:
            errors.append({"row": i, "error": "نرخ باید بزرگ‌تر از صفر باشد."})
            continue
        discount = _parse_int(row[6] if len(row) > 6 else None)
        changes.append({
            "row": i,
            "product": product.pk,
            "product_name": product.name,
            "union_name": product.union.name,
            "current_price": product.current_price,
            "new_price": new_price,
            "max_discount_percent": discount if discount is not None else product.max_discount_percent,
            "note": str(row[7]).strip() if len(row) > 7 and row[7] else "",
            "unchanged": new_price == product.current_price,
        })

    applied = 0
    if not dry_run and changes:
        with transaction.atomic():
            for c in changes:
                if c["unchanged"]:
                    continue
                set_official_price(
                    allowed[c["product"]], c["new_price"], request.user,
                    max_discount=c["max_discount_percent"],
                    note=c["note"] or "بارگذاری گروهی اداره صمت",
                )
                applied += 1

    return Response({
        "dry_run": dry_run,
        "total_rows": len(changes) + len(errors) + skipped,
        "changes": changes,
        "errors": errors,
        "skipped": skipped,
        "applied": applied,
    })
